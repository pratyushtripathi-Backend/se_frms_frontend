from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "D:/Frontend-Frms/se_frms_frontend/docs/FRMS_Audit_Page_Requirements.docx"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(9)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        set_cell_text(header_cells[index], header, True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    document.add_paragraph()


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def main():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].font.color.rgb = RGBColor(47, 84, 150)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("FRMS Audit Page Requirements")
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("UI columns, filters, detail sections, and timeline structure").italic = True

    document.add_heading("1. Purpose", level=1)
    document.add_paragraph(
        "The Audit Page is an investigation and compliance screen. It should help an admin or fraud analyst trace the full lifecycle of a transaction across services, including transaction intake, fraud evaluation, scoring, decision, event publishing, audit creation, and analytics creation."
    )

    document.add_heading("2. Main Audit Table Columns", level=1)
    add_table(
        document,
        ["Column", "Source", "Purpose"],
        [
            ["auditLogId", "se_frms_audit_log.id", "Unique audit log identifier."],
            ["transactionId", "se_frms_audit_log.transaction_id", "Used to open complete audit trail for one transaction."],
            ["serviceName", "se_frms_audit_log.service_name", "Shows which microservice generated the event."],
            ["eventType", "se_frms_audit_log.event_type", "Shows event type such as FRAUD_EVALUATION_COMPLETED."],
            ["referenceId", "se_frms_audit_log.reference_id", "Service-specific related ID, such as decisionId or scoringId."],
            ["finalDecision", "eventDetails.finalDecision", "ALLOW, REVIEW, or BLOCK decision."],
            ["totalRiskScore", "eventDetails.totalRiskScore", "Final fraud risk score."],
            ["status", "se_frms_audit_log.status", "Active/inactive audit record status."],
            ["performedBy", "se_frms_audit_log.performed_by", "Actor or service that performed the action."],
            ["createdBy", "se_frms_audit_log.created_by", "Creator of audit record, usually AUDIT_SERVICE."],
            ["createdAt", "se_frms_audit_log.created_at", "Audit record creation timestamp."],
            ["updatedAt", "se_frms_audit_log.updated_at", "Last updated timestamp."],
            ["action", "Frontend only", "View Details action."],
        ],
    )

    document.add_heading("3. Recommended Filters", level=1)
    add_table(
        document,
        ["Filter", "Type", "Backend Field"],
        [
            ["transactionId", "Text / UUID", "transaction_id"],
            ["serviceName", "Dropdown", "service_name"],
            ["eventType", "Dropdown", "event_type"],
            ["finalDecision", "Dropdown", "event_details->finalDecision"],
            ["status", "Boolean dropdown", "status"],
            ["fromDate", "Date-time", "created_at"],
            ["toDate", "Date-time", "created_at"],
            ["referenceId", "Text / UUID", "reference_id"],
        ],
    )

    document.add_heading("4. Audit Details View", level=1)
    add_bullets(
        document,
        [
            "auditLogId",
            "transactionId",
            "serviceName",
            "eventType",
            "referenceId",
            "performedBy",
            "status",
            "createdBy",
            "createdAt",
            "updatedAt",
        ],
    )

    document.add_heading("5. Event Details Section", level=1)
    add_table(
        document,
        ["Field", "Source", "Description"],
        [
            ["scoringId", "eventDetails.scoringId", "Scoring record linked to the transaction."],
            ["decisionId", "eventDetails.decisionId", "Decision record linked to the transaction."],
            ["totalRiskScore", "eventDetails.totalRiskScore", "Calculated cumulative risk score."],
            ["finalDecision", "eventDetails.finalDecision", "Final decision from Decision Service."],
            ["decisionReason", "eventDetails.decisionReason", "Reason generated from policy thresholds."],
            ["triggeredRules", "eventDetails.triggeredRules", "Matched fraud rules and their scores."],
            ["transactionData", "eventDetails.transactionData", "Full transaction payload snapshot."],
            ["eventTimestamp", "eventDetails.eventTimestamp", "Fraud event timestamp from Fraud Engine."],
        ],
    )

    document.add_heading("6. Transaction Data Section", level=1)
    add_table(
        document,
        ["Field", "Purpose"],
        [
            ["externalTransactionId", "Bank or external reference ID for idempotency."],
            ["merchantId", "Merchant identifier."],
            ["merchantName", "Merchant display name."],
            ["merchantCategory", "Merchant category such as RETAIL."],
            ["userId", "Customer/user identifier."],
            ["customerAccountNumber", "Masked customer account number."],
            ["amount", "Transaction amount."],
            ["currency", "Transaction currency."],
            ["channel", "Channel such as UPI, WEB, ATM, POS."],
            ["transactionType", "Type such as PAYMENT, TRANSFER, WITHDRAWAL."],
            ["ipAddress", "IP address used for transaction."],
            ["latitude", "Transaction latitude."],
            ["longitude", "Transaction longitude."],
            ["deviceFingerprint", "Device fingerprint value."],
            ["deviceId", "Device identifier."],
            ["os", "Device operating system."],
            ["appVersion", "Application version."],
            ["upiId", "Customer UPI ID."],
            ["beneficiaryUpiId", "Beneficiary or merchant UPI ID."],
            ["paymentMode", "Payment mode such as UPI_COLLECT."],
            ["mobile", "Customer mobile number."],
            ["location", "Human-readable location."],
        ],
    )

    document.add_heading("7. Triggered Rules Section", level=1)
    add_table(
        document,
        ["Column", "Description"],
        [
            ["ruleCode", "Unique rule code from monolithic fraud rule module."],
            ["ruleName", "Business-readable rule name."],
            ["ruleExpression", "Expression evaluated by Scoring Service."],
            ["ruleScore", "Configured score from rule score module."],
            ["calculatedScore", "Score added to totalRiskScore after match."],
            ["matched", "Whether rule matched the transaction."],
            ["categoryName", "Rule category name."],
        ],
    )

    document.add_heading("8. Transaction Timeline Section", level=1)
    add_table(
        document,
        ["Step", "Event", "Description"],
        [
            ["1", "Transaction received", "Transaction Service accepts request."],
            ["2", "Fraud evaluation started", "Transaction Service triggers Fraud Engine."],
            ["3", "Rules loaded", "Fraud Engine uses active rules from local cache."],
            ["4", "Scoring completed", "Scoring Service calculates totalRiskScore and matched rules."],
            ["5", "Decision completed", "Decision Service returns ALLOW, REVIEW, or BLOCK."],
            ["6", "Fraud event published", "Fraud Engine publishes event to Kafka."],
            ["7", "Audit record created", "Audit Service consumes event and stores audit log."],
            ["8", "Analytics record created", "Analytics Service consumes event and stores reporting record."],
        ],
    )

    document.add_heading("9. Page Layout Recommendation", level=1)
    add_bullets(
        document,
        [
            "Top area: summary counters for total audit events, blocked transactions, review transactions, and failed events.",
            "Filter bar: transactionId, serviceName, eventType, finalDecision, date range, and status.",
            "Main grid: audit table with pagination and View Details action.",
            "Details drawer/page: audit metadata, event details, transaction data, triggered rules, and timeline.",
            "Use JSON viewer only as an expandable advanced section, not as the default view.",
        ],
    )

    document.add_heading("10. API Support Needed", level=1)
    add_table(
        document,
        ["API", "Purpose"],
        [
            ["GET /api/v1/audit-logs", "Paginated audit list for main table."],
            ["GET /api/v1/audit-logs/{auditLogId}", "Open audit details by audit ID."],
            ["GET /api/v1/audit-logs/transaction/{transactionId}", "Show full timeline for one transaction."],
        ],
    )

    document.save(OUTPUT)


if __name__ == "__main__":
    main()
